"""Portable, fact-bounded DOCX rendering for the anonymous public flow.

The renderer is deliberately separate from the owner's formal-resume module:
it accepts only the public structured material payload and never reads a local
database, private profile, photo, or existing export.
"""

from __future__ import annotations

from io import BytesIO
import re
from typing import Any

from PIL import Image, UnidentifiedImageError
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BLUE = RGBColor(46, 116, 181)
NAVY = RGBColor(24, 55, 86)
INK = RGBColor(34, 43, 54)
MUTED = RGBColor(91, 108, 124)
PALE_BLUE = "EAF3FA"
FONT_CN = "Microsoft YaHei"

TEMPLATE_TOKENS = {
    "{{姓名}}", "{{目标岗位}}", "{{联系方式}}", "{{职业概述}}", "{{核心技能}}",
    "{{工作经历}}", "{{项目经历}}", "{{工程作品}}", "{{教育背景}}", "{{证书}}", "{{岗位亮点}}", "{{投递前核对}}",
}


def render_public_resume_docx(materials: dict[str, Any], photo_bytes: bytes | None = None) -> bytes:
    """Build the generic, editable Engineering Professional resume.

    This is intentionally a neutral default rather than a copy of any
    visitor's private résumé.  The fixed hierarchy is suitable for common
    Chinese engineering roles, while all substantive wording comes from the
    confirmed public profile selected for the current JD.
    """
    resume = materials["resume"]
    readiness = materials.get("export_readiness") or {}
    if not readiness.get("ready"):
        raise ValueError("正式简历需要通过资料完整性检查后才能生成")
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(0.86), Cm(0.78)
    section.left_margin = section.right_margin = Cm(1.15)
    section.header_distance = section.footer_distance = Cm(0.45)
    _configure(doc)

    identity = resume["identity"]
    _identity_header(doc, identity, resume.get("target_title") or "目标岗位", photo_bytes)

    _summary(doc, resume.get("summary", ""))
    _section(doc, "专业能力")
    _skills(doc, resume.get("skills", []))
    _entries(doc, "工作经历", resume.get("experiences", []), "company", "role")
    _entries(doc, "项目经历", resume.get("projects", []), "name", "subtitle")
    _portfolio(doc, resume.get("portfolio", []))

    education = resume.get("education") or {}
    if education:
        _section(doc, "教育背景与专业资质")
        _labeled(doc, "教育", " | ".join(item for item in (education.get("school", ""), education.get("major", ""), education.get("degree", ""), education.get("period", "")) if item))
    credentials = [item.get("name", "") for item in resume.get("credentials", []) if item.get("name")]
    if credentials:
        if not education:
            _section(doc, "教育背景与专业资质")
        _labeled(doc, "证书", " · ".join(credentials))
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def normalize_public_resume_photo(photo_bytes: bytes) -> bytes:
    """Validate and center-crop a visitor photo to a compact 3:4 JPEG.

    The image stays in request memory.  Restricting formats and dimensions
    makes the generated DOCX predictable and avoids accepting disguised files.
    """
    if not photo_bytes or len(photo_bytes) > 2 * 1024 * 1024:
        raise ValueError("证件照不能为空且不能超过 2MB")
    try:
        with Image.open(BytesIO(photo_bytes)) as source:
            if source.format not in {"JPEG", "PNG"}:
                raise ValueError("证件照仅支持 JPG 或 PNG")
            width, height = source.size
            if width < 160 or height < 200 or width > 6000 or height > 6000:
                raise ValueError("证件照尺寸需介于 160×200 与 6000×6000 像素之间")
            image = source.convert("RGB")
    except UnidentifiedImageError as error:
        raise ValueError("无法识别证件照图片") from error
    target_ratio = 3 / 4
    source_ratio = image.width / image.height
    if source_ratio > target_ratio:
        crop_width = round(image.height * target_ratio)
        left = (image.width - crop_width) // 2
        image = image.crop((left, 0, left + crop_width, image.height))
    else:
        crop_height = round(image.width / target_ratio)
        top = (image.height - crop_height) // 2
        image = image.crop((0, top, image.width, top + crop_height))
    image.thumbnail((360, 480), Image.Resampling.LANCZOS)
    output = BytesIO()
    image.save(output, format="JPEG", quality=88, optimize=True)
    return output.getvalue()


def _identity_header(doc: Document, identity: dict[str, Any], title: str, photo_bytes: bytes | None) -> None:
    name = identity.get("name") or "候选人"
    contact = " · ".join(item for item in (identity.get("location", ""), identity.get("phone", ""), identity.get("email", "")) if item)
    if not photo_bytes:
        heading = doc.add_paragraph()
        heading.paragraph_format.space_after = Pt(1)
        _run(heading, name, 23, NAVY, True)
        role = doc.add_paragraph()
        role.paragraph_format.space_after = Pt(3)
        _run(role, title, 12.2, BLUE, True)
        if identity.get("tagline"):
            _run(role, f"  |  {identity['tagline']}", 8.7, MUTED)
        sub = doc.add_paragraph()
        sub.paragraph_format.space_after = Pt(0)
        _run(sub, contact, 8.7, INK)
        return

    photo = normalize_public_resume_photo(photo_bytes)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    _remove_table_borders(table)
    # ``cell.width`` alone does not update a fixed table's grid in python-docx.
    # Word then keeps two equal columns and can clip the identity column when a
    # photo is present. Set both the grid and the cells so the left header text
    # and the right 3:4 photo consistently share the available page width.
    table.columns[0].width = Cm(15.85)
    table.columns[1].width = Cm(2.6)
    left, right = table.rows[0].cells
    left.width, right.width = Cm(15.85), Cm(2.6)
    heading = left.paragraphs[0]
    heading.paragraph_format.space_after = Pt(1)
    _run(heading, name, 23, NAVY, True)
    role = left.add_paragraph()
    role.paragraph_format.space_after = Pt(3)
    _run(role, title, 12.2, BLUE, True)
    if identity.get("tagline"):
        _run(role, f"  |  {identity['tagline']}", 8.7, MUTED)
    sub = left.add_paragraph()
    sub.paragraph_format.space_after = Pt(0)
    _run(sub, contact, 8.7, INK)
    picture = right.paragraphs[0]
    picture.alignment = 2
    picture.paragraph_format.space_after = Pt(0)
    picture.add_run().add_picture(BytesIO(photo), width=Cm(2.28), height=Cm(3.04))


def _remove_table_borders(table: Any) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "nil")


def render_public_resume_from_template_docx(materials: dict[str, Any], template_bytes: bytes) -> bytes:
    """Fill a visitor-supplied DOCX template entirely in memory.

    A custom template is an opt-in layout mechanism, not a promise that every
    arbitrary Word file can be reconstructed.  The supported placeholders
    must be unbroken text in one Word run; they preserve that run's style.
    """
    readiness = materials.get("export_readiness") or {}
    if not readiness.get("ready"):
        raise ValueError("正式简历需要通过资料完整性检查后才能生成")
    if not template_bytes or len(template_bytes) > 2 * 1024 * 1024:
        raise ValueError("DOCX 模板不能为空且不能超过 2MB")
    try:
        doc = Document(BytesIO(template_bytes))
    except Exception as error:  # python-docx exposes several zip/XML errors.
        raise ValueError("无法识别该 DOCX 模板") from error

    context = _template_context(materials["resume"])
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        paragraphs.extend(_table_paragraphs(table))
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
        for table in section.header.tables:
            paragraphs.extend(_table_paragraphs(table))
        for table in section.footer.tables:
            paragraphs.extend(_table_paragraphs(table))

    all_text = "\n".join(paragraph.text for paragraph in paragraphs)
    placeholders = set(re.findall(r"\{\{[^{}]+\}\}", all_text))
    unknown = placeholders - TEMPLATE_TOKENS
    if unknown:
        raise ValueError(f"模板含有不支持的占位符：{'、'.join(sorted(unknown))}")
    required = {"{{姓名}}", "{{目标岗位}}", "{{核心技能}}"}
    if not required.issubset(placeholders):
        raise ValueError("DOCX 模板至少需要包含 {{姓名}}、{{目标岗位}}、{{核心技能}}")

    for paragraph in paragraphs:
        _replace_template_tokens(paragraph, context)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _configure(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        normal._element.rPr.rFonts.set(qn(key), FONT_CN)
    normal.font.size = Pt(9)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(1.25)
    normal.paragraph_format.line_spacing = 1.04
    bullet = doc.styles["List Bullet"]
    bullet.font.name = FONT_CN
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        bullet._element.rPr.rFonts.set(qn(key), FONT_CN)
    bullet.font.size = Pt(8.6)
    bullet.font.color.rgb = INK
    bullet.paragraph_format.left_indent = Cm(0.42)
    bullet.paragraph_format.first_line_indent = Cm(-0.22)
    bullet.paragraph_format.space_before = Pt(0)
    bullet.paragraph_format.space_after = Pt(1.15)
    bullet.paragraph_format.line_spacing = 1.03


def _section(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    _shade(paragraph, "2E74B5")
    _run(paragraph, f"  {text}", 10.3, RGBColor(255, 255, 255), True)


def _skills(doc: Document, groups: list[dict[str, Any]]) -> None:
    for group in groups:
        _labeled(doc, str(group.get("label", "核心技能")), " · ".join(group.get("items", [])))


def _entries(doc: Document, heading: str, entries: list[dict[str, Any]], first: str, second: str) -> None:
    if not entries:
        return
    _section(doc, heading)
    for entry in entries:
        title = doc.add_paragraph()
        title.paragraph_format.space_before = Pt(1.3)
        title.paragraph_format.space_after = Pt(1.3)
        title.paragraph_format.keep_with_next = True
        _run(title, str(entry.get(first, "")), 9.7, NAVY, True)
        if entry.get(second):
            _run(title, f"  |  {entry[second]}", 9.15, INK, True)
        if entry.get("period"):
            _run(title, f"  |  {entry['period']}", 8.95, MUTED)
        for bullet in entry.get("bullets", []):
            paragraph = doc.add_paragraph(style="List Bullet")
            _run(paragraph, bullet.get("text", ""), 9.2, INK)


def _summary(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2.2)
    paragraph.paragraph_format.space_after = Pt(2.2)
    paragraph.paragraph_format.left_indent = Cm(0.22)
    paragraph.paragraph_format.right_indent = Cm(0.22)
    _shade(paragraph, PALE_BLUE)
    _run(paragraph, "职业概况  ", 8.8, BLUE, True)
    _run(paragraph, text, 9.35, INK)


def _portfolio(doc: Document, entries: list[dict[str, Any]]) -> None:
    if not entries:
        return
    _section(doc, "工程作品")
    for entry in entries:
        heading = str(entry.get("name", ""))
        if entry.get("period"):
            heading += f"  |  {entry['period']}"
        _labeled(doc, heading, str(entry.get("summary", "")))


def _labeled(doc: Document, label: str, value: str) -> None:
    if not value:
        return
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1.2)
    _run(paragraph, f"{label}：", 9.15, NAVY, True)
    _run(paragraph, value, 9.15, INK)


def _run(paragraph: Any, text: str, size: float, color: RGBColor, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    run.font.name = FONT_CN
    fonts = run._element.get_or_add_rPr().rFonts
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        fonts.set(qn(key), FONT_CN)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def _shade(paragraph: Any, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _table_paragraphs(table: Any) -> list[Any]:
    paragraphs: list[Any] = []
    for row in table.rows:
        for cell in row.cells:
            paragraphs.extend(cell.paragraphs)
            for nested in cell.tables:
                paragraphs.extend(_table_paragraphs(nested))
    return paragraphs


def _template_context(resume: dict[str, Any]) -> dict[str, str]:
    identity = resume.get("identity") or {}
    contact = " · ".join(item for item in (identity.get("location", ""), identity.get("phone", ""), identity.get("email", "")) if item)
    skills = "\n".join(f"{item.get('label', '核心技能')}：{' · '.join(item.get('items', []))}" for item in resume.get("skills", []))
    work = _template_entries(resume.get("experiences", []), "company", "role")
    projects = _template_entries(resume.get("projects", []), "name", "subtitle")
    portfolio = "\n\n".join("\n".join(part for part in (" · ".join(value for value in (item.get("name", ""), item.get("period", "")) if value), item.get("summary", "")) if part) for item in resume.get("portfolio", []))
    education = resume.get("education") or {}
    education_text = " · ".join(item for item in (education.get("school", ""), education.get("major", ""), education.get("degree", ""), education.get("period", "")) if item)
    credentials = "、".join(item.get("name", "") for item in resume.get("credentials", []) if item.get("name"))
    highlights = "；".join(
        f"{item.get('label', '核心技能')}：{'、'.join(item.get('items', [])[:4])}"
        for item in resume.get("skills", [])[:2]
        if item.get("items")
    )
    return {
        "{{姓名}}": str(identity.get("name", "")),
        "{{目标岗位}}": str(resume.get("target_title", "")),
        "{{联系方式}}": contact,
        "{{职业概述}}": str(resume.get("summary", "")),
        "{{核心技能}}": skills,
        "{{工作经历}}": work,
        "{{项目经历}}": projects,
        "{{工程作品}}": portfolio,
        "{{教育背景}}": education_text,
        "{{证书}}": credentials,
        "{{岗位亮点}}": highlights,
        # Compatibility only: old templates may still carry this token.  Internal
        # checking instructions must never be inserted into an external resume.
        "{{投递前核对}}": "",
    }


def _template_entries(entries: list[dict[str, Any]], first: str, second: str) -> str:
    blocks = []
    for entry in entries:
        heading = " · ".join(item for item in (entry.get(first, ""), entry.get(second, ""), entry.get("period", "")) if item)
        bullets = "\n".join(f"• {item.get('text', '')}" for item in entry.get("bullets", []) if item.get("text"))
        blocks.append("\n".join(part for part in (heading, bullets) if part))
    return "\n\n".join(blocks)


def _replace_template_tokens(paragraph: Any, context: dict[str, str]) -> None:
    combined = "".join(run.text for run in paragraph.runs)
    for token in TEMPLATE_TOKENS:
        if token not in combined:
            continue
        matching_runs = [run for run in paragraph.runs if token in run.text]
        if not matching_runs:
            raise ValueError(f"占位符 {token} 不能拆分为多段文字；请在 Word 中重新输入该占位符")
        for run in matching_runs:
            run.text = run.text.replace(token, context[token])
