from io import BytesIO
import importlib.util
from pathlib import Path

import pytest
from docx import Document
from pydantic import ValidationError

from app.public_product.demo_profiles import public_demo_profiles
from app.public_product.documents import render_public_resume_docx
from app.public_product.materials import compose_public_materials
from app.public_product.profile_contract import (
    PublicJobInput,
    PublicProfile,
    PublicProfileAnalysisRequest,
    analyze_public_profile,
    assess_export_readiness,
)


def test_demo_profiles_are_fictional_and_export_ready():
    profiles = public_demo_profiles()
    assert len(profiles) == 3
    assert all(profile.origin == "sample" for profile in profiles)
    assert all(assess_export_readiness(profile)["ready"] for profile in profiles)
    assert all(not profile.identity.phone and not profile.identity.email for profile in profiles)


def test_profile_contract_rejects_unknown_fields():
    with pytest.raises(ValidationError):
        PublicProfile.model_validate({"unknown": "not allowed"})


def test_analysis_and_materials_use_confirmed_demo_evidence():
    profile = public_demo_profiles()[0]
    job = PublicJobInput(
        title="PLC 电气工程师",
        description="负责 PLC、HMI、Modbus 通信与现场异常排查。",
        location="深圳",
    )
    request = PublicProfileAnalysisRequest(profile=profile, job=job)
    analysis = analyze_public_profile(request)
    materials = compose_public_materials(profile, job)
    assert analysis["export_readiness"]["ready"]
    assert materials["review"]["hr_review"]
    assert materials["review"]["technical_review"]
    assert "PLC" in materials["ats_text"]


def test_incomplete_profile_cannot_export_placeholders():
    materials = compose_public_materials(
        PublicProfile(),
        PublicJobInput(title="测试岗位", description="测试岗位要求"),
    )
    with pytest.raises(ValueError, match="资料完整性检查"):
        render_public_resume_docx(materials)


def test_docx_is_editable_and_excludes_internal_review():
    profile = public_demo_profiles()[0]
    materials = compose_public_materials(
        profile,
        PublicJobInput(title="PLC 电气工程师", description="PLC、HMI、Modbus 与现场调试。"),
    )
    document = Document(BytesIO(render_public_resume_docx(materials)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "PLC 电气工程师" in text
    assert "HR 与技术主管" not in text
    assert "投递前核对" not in text


def test_anonymous_http_function_returns_health_and_docx():
    module_path = Path("deployment/zaotu-function/main.py")
    spec = importlib.util.spec_from_file_location("zaotu_public_function_test", module_path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    profile = public_demo_profiles()[0]
    request = PublicProfileAnalysisRequest(
        profile=profile,
        job=PublicJobInput(title="PLC 电气工程师", description="PLC、HMI、Modbus 与现场调试。"),
    )
    client = module.app.test_client()
    health = client.get("/api/healthz")
    document = client.post(
        "/api/public/guest/profile/materials/docx",
        json=request.model_dump(mode="json"),
    )
    assert health.status_code == 200
    assert health.headers["Content-Disposition"] == "inline"
    assert document.status_code == 200
    assert document.data.startswith(b"PK")
    assert document.headers["Content-Disposition"].startswith("attachment;")
